import re
from datetime import datetime
from io import BytesIO
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from apps.modules.import_export import ImportFileError, ParsedRecord, parse_file

EXPECTED_AMOUNT = 100


def workbook_bytes(rows: list[list[Any]] | None = None, **options: Any) -> bytes:
    """Создать книгу с настраиваемыми данными и структурой для контрактных проверок."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = options.get("title", "Расходы")
    if options.get("mac"):
        workbook.epoch = CALENDAR_MAC_1904
    sheet.append(["Дата", "Сумма", "Категория", "Описание"])
    for row in rows or []:
        sheet.append(row)
    if options.get("extra_sheet"):
        workbook.create_sheet("Лишний")
    if options.get("hidden_row"):
        sheet.row_dimensions[2].hidden = True
    if options.get("hidden_column"):
        sheet.column_dimensions["E"].hidden = True
    if options.get("merged"):
        sheet.merge_cells("C2:D2")
    if options.get("styled_empty"):
        sheet["Z1000000"].number_format = "0"
    if options.get("string_formula"):
        sheet["D2"].data_type = "s"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def parse_xlsx(content: bytes, content_type: str | None = None) -> list[ParsedRecord]:
    """Разобрать XLSX через публичный API модуля."""
    return parse_file(
        content, transaction_type="outcome", file_format="xlsx", filename="test.XLSX", content_type=content_type
    )


def xlsx_error(content: bytes, code: str = "FM-422301") -> ImportFileError:
    """Убедиться в контрактном коде ошибки XLSX."""
    with pytest.raises(ImportFileError) as caught:
        parse_xlsx(content)
    assert caught.value.code == code
    return caught.value


def replace_part(content: bytes, name: str, old: bytes, new: bytes) -> bytes:
    """Заменить фрагмент OOXML, сохранив корректный ZIP-контейнер."""
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            target.writestr(item.filename, data.replace(old, new) if item.filename == name else data)
    return output.getvalue()


@pytest.mark.parametrize("mac", [False, True])
@pytest.mark.parametrize("date_value", ["01.10.2026", "01/10/2026", "01-10-2026", datetime.fromisoformat("2026-10-01")])
def test_xlsx_dates(date_value: object, mac: bool) -> None:
    """Текстовые и нативные даты обеих систем Excel дают одну и ту же полночь."""
    records = parse_xlsx(workbook_bytes([[date_value, 100, " Еда ", None]], mac=mac))
    assert records[0].transaction_date == datetime.fromisoformat("2026-10-01")
    assert records[0].money_sum == EXPECTED_AMOUNT
    assert records[0].category_name == "Еда"
    assert records[0].description is None


@pytest.mark.parametrize(
    ("row", "field"),
    [
        ([datetime.fromisoformat("2026-10-01T00:00:01"), 100, "Еда", ""], "date"),
        ([45000, 100, "Еда", ""], "date"),
        ([True, 100, "Еда", ""], "date"),
        (["01.10.2026", True, "Еда", ""], "amount"),
        (["01.10.2026", 1.01, "Еда", ""], "amount"),
        (["01.10.2026", "100.0", "Еда", ""], "amount"),
        (["01.10.2026", 1, 123, ""], "category"),
        (["01.10.2026", 1, "Еда", True], "description"),
        (["01.10.2026", 1, "Еда", "#DIV/0!"], "description"),
    ],
)
def test_xlsx_invalid_types(row: list[Any], field: str) -> None:
    """Типы ячеек не приводятся молча к строкам и целым числам."""
    exc = xlsx_error(workbook_bytes([row]), "FM-422303")
    assert [(issue.row, issue.field, issue.reason) for issue in exc.errors] == [(2, field, f"invalid_{field}")]


@pytest.mark.parametrize("column", range(4))
def test_xlsx_formula(column: int) -> None:
    """Формула в каждой колонке имеет отдельную ошибку, даже с кешированным результатом."""
    row: list[Any] = ["01.10.2026", 1, "Еда", ""]
    row[column] = "=1+1"
    data = replace_part(workbook_bytes([row]), "xl/worksheets/sheet1.xml", b"<f>1+1</f><v></v>", b"<f>1+1</f><v>2</v>")
    exc = xlsx_error(data, "FM-422303")
    assert exc.errors[0].reason == "formula_not_allowed"
    assert exc.errors[0].field == ["date", "amount", "category", "description"][column]


def test_literal_formula() -> None:
    """Строковая ячейка со знаком равенства остаётся текстом."""
    records = parse_xlsx(workbook_bytes([["01.10.2026", 1, "Еда", "=пояснение"]], string_formula=True))
    assert records[0].description == "=пояснение"


@pytest.mark.parametrize(
    "options",
    [{"title": "Доходы"}, {"extra_sheet": True}, {"hidden_row": True}, {"hidden_column": True}, {"merged": True}],
)
def test_xlsx_structure(options: dict[str, Any]) -> None:
    """Лишние листы, скрытые области, объединения и неправильное имя запрещены."""
    xlsx_error(workbook_bytes([["01.10.2026", 1, "Еда", ""]], **options))


def test_xlsx_sparse_rows() -> None:
    """Пустые строки не теряют номера; форматирование далёкой ячейки не является данными."""
    records = parse_xlsx(workbook_bytes([[], ["01.10.2026", 1, "Еда", ""]], styled_empty=True))
    assert [record.row for record in records] == [3]


def test_xlsx_extra_column() -> None:
    """Данные за D запрещены, даже если первые четыре колонки пусты."""
    xlsx_error(workbook_bytes([[None, None, None, None, "=1+1"]]))


def test_exact_numeric_xml() -> None:
    """Дробь меньше точности float остаётся ошибкой суммы."""
    data = replace_part(
        workbook_bytes([["01.10.2026", 100, "Еда", ""]]),
        "xl/worksheets/sheet1.xml",
        b"<v>100</v>",
        b"<v>100.00000000000000001</v>",
    )
    assert xlsx_error(data, "FM-422303").errors[0].reason == "invalid_amount"


def test_submillisecond_date() -> None:
    """Даже малая ненулевая доля суток не округляется до полуночи."""
    data = workbook_bytes([[datetime.fromisoformat("2026-10-01"), 1, "Еда", ""]])
    with ZipFile(BytesIO(data)) as source:
        xml = source.read("xl/worksheets/sheet1.xml")
    number = re.search(rb'<c r="A2"[^>]*><v>([0-9]+)</v>', xml)
    assert number is not None
    data = replace_part(data, "xl/worksheets/sheet1.xml", number[1], number[1] + b".000000000001")
    assert xlsx_error(data, "FM-422303").errors[0].reason == "invalid_date"


@pytest.mark.parametrize("content", [b"PKbroken", b"PK\x03\x04"])
def test_broken_zip(content: bytes) -> None:
    """Повреждённый ZIP даёт структурную ошибку, а не необработанное исключение."""
    xlsx_error(content)


def test_non_zip() -> None:
    """Произвольные байты вместо XLSX дают неподдерживаемый формат."""
    xlsx_error(b"not a zip", "FM-415301")


def test_xml_entities() -> None:
    """DTD и сущности не разрешаются при чтении OOXML."""
    data = replace_part(
        workbook_bytes(), "xl/workbook.xml", b"<workbook ", b'<!DOCTYPE workbook [<!ENTITY x "data">]><workbook '
    )
    xlsx_error(data)


def test_external_relationship() -> None:
    """Внешние связи не загружаются и запрещены контрактом."""
    data = replace_part(workbook_bytes(), "xl/_rels/workbook.xml.rels", b"Target=", b'TargetMode="External" Target=')
    xlsx_error(data)


def test_zip_limits() -> None:
    """Число ZIP-элементов и распакованный объём ограничены до разбора XML."""
    for count, payload in [(1001, b""), (1, b"0" * (50 * 1024 * 1024 + 1))]:
        stream = BytesIO()
        with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
            for index in range(count):
                archive.writestr(f"item{index}", payload)
        xlsx_error(stream.getvalue(), "FM-413302")


def test_ooxml_escaped_controls() -> None:
    """Экранированные OOXML символы проверяются после декодирования и нормализации CR."""
    data = workbook_bytes([["01.10.2026", 1, "Еда", "строка_x000D_\nещё"]])
    assert parse_xlsx(data)[0].description == "строка\nещё"
    data = workbook_bytes([["01.10.2026", 1, "Еда", "_x0000_"]])
    assert xlsx_error(data, "FM-422303").errors[0].reason == "invalid_description"


def test_macro_relationship() -> None:
    """Запрещённая VBA-связь обнаруживается независимо от имени целевого файла."""
    data = replace_part(
        workbook_bytes([["01.10.2026", 1, "Еда", ""]]), "xl/_rels/workbook.xml.rels", b'/styles"', b'/vbaProject"'
    )
    xlsx_error(data)


def test_unknown_xml_encoding() -> None:
    """Неизвестная XML-кодировка возвращает ошибку структуры вместо LookupError."""
    data = replace_part(
        workbook_bytes(),
        "xl/worksheets/sheet1.xml",
        b"<worksheet ",
        b'<?xml version="1.0" encoding="unknown"?><worksheet ',
    )
    exc = xlsx_error(data)
    assert exc.errors[0].reason == "invalid_structure"


@pytest.mark.parametrize("amount", [1, 100, 1000000000])
def test_amount_with_date_display_format(amount: int) -> None:
    """Формат отображения даты в B не изменяет точное числовое значение суммы."""
    data = workbook_bytes([[datetime.fromisoformat("2026-10-01"), amount, "Еда", ""]])
    data = replace_part(data, "xl/worksheets/sheet1.xml", b'r="B2" t="n"', b'r="B2" s="1" t="n"')
    assert parse_xlsx(data)[0].money_sum == amount
