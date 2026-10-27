from http import HTTPStatus
from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from tests.conftest import TokenFactory


@pytest.mark.parametrize("transaction_type", ["income", "outcome"])
@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
async def test_template(client: AsyncClient, make_token: TokenFactory, transaction_type: str, file_format: str) -> None:
    """Шаблон содержит только заголовок и доступен браузеру после авторизации."""
    response = await client.get(
        f"/import-export/templates/{transaction_type}?format={file_format}",
        headers={"Authorization": f"Bearer {make_token()}", "Origin": "http://localhost:5173"},
    )
    assert response.status_code == HTTPStatus.OK
    assert (
        response.headers["content-disposition"] == f'attachment; filename="{transaction_type}-template.{file_format}"'
    )
    assert "content-disposition" in response.headers["access-control-expose-headers"].lower()
    if file_format == "csv":
        assert response.headers["content-type"] == "text/csv; charset=utf-8"
        assert response.content == "\ufeffДата;Сумма;Категория;Описание\r\n".encode()
    else:
        assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.sheetnames == [("Доходы" if transaction_type == "income" else "Расходы")]
        sheet = workbook.active
        assert sheet is not None
        assert list(sheet.values) == [("Дата", "Сумма", "Категория", "Описание")]
        assert [sheet.column_dimensions[col].number_format for col in "ABCD"] == ["@", "0", "@", "@"]


async def test_template_requires_authentication(client: AsyncClient) -> None:
    """Скачивание без JWT отклоняется."""
    response = await client.get("/import-export/templates/outcome?format=csv")
    assert response.status_code == HTTPStatus.UNAUTHORIZED


@pytest.mark.parametrize("suffix", ["outcome", "other?format=csv", "income?format=xls"])
async def test_template_requires_valid_parameters(client: AsyncClient, make_token: TokenFactory, suffix: str) -> None:
    """Тип и формат обязательны и ограничены контрактом."""
    response = await client.get(
        f"/import-export/templates/{suffix}", headers={"Authorization": f"Bearer {make_token()}"}
    )
    assert response.status_code == HTTPStatus.UNPROCESSABLE_ENTITY
    assert response.json()["code"] == "FM-422000"
